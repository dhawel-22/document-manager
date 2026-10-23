"""Test that the Word-reading library works (spec/summaries.md, step 2).

The library is python-docx. This test makes one fake Word document in a
temporary folder, reads it back, and checks that the title, the heading
and the text come out exactly as written.
"""

import docx


def test_reads_a_fake_word_document(tmp_path):
    fake = tmp_path / "fake.docx"
    made = docx.Document()
    made.core_properties.title = "Fake report"
    made.add_heading("Section one", level=1)
    made.add_paragraph("This is a fake document for the tests.")
    made.save(fake)

    read = docx.Document(fake)
    assert read.core_properties.title == "Fake report"
    paragraphs = [(p.style.name, p.text) for p in read.paragraphs]
    assert paragraphs == [
        ("Heading 1", "Section one"),
        ("Normal", "This is a fake document for the tests."),
    ]
