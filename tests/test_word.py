"""Tests for one document's summary, spec/summaries.md v1.1 (SM-3, SM-5, SM-6).

Every Word document here is a fake, made by the test in a temporary folder.
"""

import re
import zipfile

import docx
import pytest

from document_manager.word import CouldNotRead, Summary, summarize_document

OLD_FORMAT = bytes.fromhex("D0CF11E0A1B11AE1") + bytes(504)


def fake_docx(path, title=None, blocks=()):
    """Make a fake Word document. blocks are (kind, text) in order, where kind
    is "title", "h1" to "h9", or "p" for a plain paragraph."""
    made = docx.Document()
    if title is not None:
        made.core_properties.title = title
    for kind, text in blocks:
        if kind == "title":
            made.add_heading(text, level=0)
        elif kind == "p":
            made.add_paragraph(text)
        else:
            made.add_heading(text, level=int(kind[1:]))
    made.save(path)
    return path


def without_saved_properties(source, target):
    """Copy a fake .docx without its saved properties, as some programs write them."""
    with zipfile.ZipFile(source) as old, zipfile.ZipFile(target, "w") as new:
        for item in old.infolist():
            data = old.read(item.filename)
            if item.filename == "docProps/core.xml":
                continue
            if item.filename == "_rels/.rels":
                data = re.sub(rb"<Relationship [^>]*core-properties[^>]*/>", b"", data)
            if item.filename == "[Content_Types].xml":
                data = re.sub(rb"<Override [^>]*/docProps/core\.xml[^>]*/>", b"", data)
            new.writestr(item, data)
    return target


def test_sm3_title_headings_and_words(tmp_path):
    path = fake_docx(tmp_path / "report.docx", title="Weekly report", blocks=[
        ("h1", "Summary"),
        ("p", "Markets were calm this week."),
        ("h2", "Details"),
        ("p", "Nothing else to add."),
    ])
    assert summarize_document(path) == Summary(
        title="Weekly report",
        headings=[(1, "Summary"), (2, "Details")],
        first_words="Summary Markets were calm this week. Details Nothing else to add.",
        word_count=11,
    )


def test_sm3_only_the_first_60_words(tmp_path):
    text = " ".join(f"word{n}" for n in range(1, 101))
    path = fake_docx(tmp_path / "long.docx", title="Long", blocks=[("p", text)])
    summary = summarize_document(path)
    assert summary.first_words == " ".join(f"word{n}" for n in range(1, 61))
    assert summary.word_count == 100


def test_sm3_title_from_the_title_paragraph(tmp_path):
    path = fake_docx(tmp_path / "t.docx", title="", blocks=[
        ("title", "Risk Watch"),
        ("p", "Text."),
    ])
    assert summarize_document(path).title == "Risk Watch"


def test_sm3_no_title_is_never_made_up(tmp_path):
    plain = fake_docx(tmp_path / "plain.docx", blocks=[("p", "Just text.")])
    bare = without_saved_properties(plain, tmp_path / "bare.docx")
    assert "docProps/core.xml" not in zipfile.ZipFile(bare).namelist()
    # The library's usual way would make up a title here:
    assert docx.Document(str(bare)).core_properties.title == "Word Document"
    assert summarize_document(bare).title == ""


def test_sm3_tables_read_in_order_and_merged_cells_once(tmp_path):
    made = docx.Document()
    made.core_properties.title = "Table"
    made.add_paragraph("Before.")
    table = made.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged header"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    made.add_paragraph("After.")
    path = tmp_path / "table.docx"
    made.save(path)
    summary = summarize_document(path)
    assert summary.first_words == "Before. Merged header A B After."
    assert summary.word_count == 6


def test_sm3_empty_document(tmp_path):
    path = fake_docx(tmp_path / "empty.docx", title="")
    assert summarize_document(path) == Summary(
        title="", headings=[], first_words="", word_count=0
    )


def test_sm5_only_reads(tmp_path):
    path = fake_docx(tmp_path / "report.docx", title="Weekly report", blocks=[
        ("p", "Text."),
    ])
    before = (path.read_bytes(), path.stat().st_mtime_ns)
    summarize_document(path)
    assert (path.read_bytes(), path.stat().st_mtime_ns) == before


def test_sm6_old_word_file(tmp_path):
    old = tmp_path / "old.doc"
    old.write_bytes(OLD_FORMAT)
    with pytest.raises(CouldNotRead, match=r"old Word file \(\.doc\)"):
        summarize_document(old)


def test_sm6_protected_file(tmp_path):
    protected = tmp_path / "secret.docx"
    protected.write_bytes(OLD_FORMAT)
    with pytest.raises(CouldNotRead, match="protected with a password"):
        summarize_document(protected)


def test_sm6_word_lock_file(tmp_path):
    lock = tmp_path / "~$report.docx"
    lock.write_bytes(b"\x05Ernie" + bytes(150))
    with pytest.raises(CouldNotRead, match="lock file"):
        summarize_document(lock)


def test_sm6_broken_file_gives_the_error_word_for_word(tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"this is not a Word document")
    with pytest.raises(CouldNotRead, match="PackageNotFoundError"):
        summarize_document(broken)


def test_sm6_missing_file(tmp_path):
    with pytest.raises(CouldNotRead, match="FileNotFoundError"):
        summarize_document(tmp_path / "gone.docx")
