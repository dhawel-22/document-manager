"""Tests for long runs, spec/long-runs.md (LR-1 to LR-9): Register and Check
keep what they've read in a progress file, and the next run picks up where
the last one stopped. The window's words for it are tested with the window.

Every folder here is fake, made by the test in a temporary folder. A run is
stopped by a counter, as the window's Stop does, by a problem, by Ctrl+C,
or by ending its process at once, as a power cut would.
"""

import ctypes
import re
import sqlite3
import subprocess
import sys
import time
from contextlib import contextmanager
from ctypes import wintypes
from pathlib import Path

import docx
import pytest

from document_manager.__main__ import main
from document_manager.accept import accept_folder
from document_manager.check import CheckResult, check_folder
from document_manager.commands import accept, changes_to_accept, run
from document_manager.fingerprint import fingerprint, fingerprint_of
from document_manager.progress import Progress
from document_manager.register import IN_USE, read_file, register_folder

PROJECT = Path(__file__).resolve().parent.parent
PLACES = ["letters/note.txt", "letters/plan.txt", "mail/archive.pst", "report.txt"]
# A counter is told 1, 3 and 4 files found (the top folder, then letters,
# then mail), then 0 to 4 read.
FOUND = [(1, None), (3, None), (4, None)]

# Ends a register command at once, as a power cut would, once it has kept two
# files: it keeps after every file, so the test needn't wait a minute, and it
# waits while it reads the third.
POWER_CUT = """
import sys, time
import document_manager.progress as progress
import document_manager.register as register
from document_manager.__main__ import main

progress.KEEP_EVERY = 0
real, read = register.fingerprint_of, []

def slow(f):
    read.append(f)
    if len(read) == 3:
        time.sleep(120)
    return real(f)

register.fingerprint_of = slow
main(sys.argv[1:])
"""

kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
kernel32.CreateFileW.restype = wintypes.HANDLE
kernel32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                 wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD,
                                 wintypes.HANDLE]
kernel32.CloseHandle.argtypes = [wintypes.HANDLE]


class Stopped(Exception):
    """What a counter raises to stop a run, as the window's Stop does."""


@contextmanager
def held_open(path):
    """Hold a file open with no sharing, the way Outlook holds its mail files."""
    handle = kernel32.CreateFileW(str(path), 0x80000000, 0, None, 3, 0, None)
    assert handle != wintypes.HANDLE(-1).value, f"Windows error {ctypes.get_last_error()}"
    try:
        yield
    finally:
        kernel32.CloseHandle(handle)


def make_docs(tmp_path):
    """A small folder of fake files: one at the top, three in subfolders."""
    docs = tmp_path / "docs"
    (docs / "letters").mkdir(parents=True)
    (docs / "mail").mkdir()
    for place in PLACES:
        (docs / place).write_bytes(f"fake {place}\n".encode())
    return docs


def stops_after(files):
    """A counter that stops the run once it has read this many files."""

    def counter(done, total):
        if total is not None and done == files:
            raise Stopped

    return counter


def progress_of(register_file):
    return register_file.with_name(register_file.name + ".progress")


def query(database, sql):
    con = sqlite3.connect(f"{database.as_uri()}?mode=ro", uri=True)
    try:
        return con.execute(sql).fetchall()
    finally:
        con.close()


def kept_places(register_file):
    """The places the progress file next to the register file has kept."""
    progress = progress_of(register_file)
    if not progress.exists():
        return []
    try:
        return [place for (place,) in query(progress, "SELECT path FROM file ORDER BY path")]
    except sqlite3.Error:  # not made yet, or being written right now
        return []


def the_list(register_file):
    """Every file the register file lists, with its size and fingerprint."""
    return query(register_file, "SELECT path, size, fingerprint FROM file ORDER BY path")


def reads(monkeypatch):
    """Note the name of every file a run opens to read."""
    opened = []

    def noting(path):
        opened.append(path.name)
        return read_file(path)

    monkeypatch.setattr("document_manager.register.read_file", noting)
    return opened


def test_lr1_a_stopped_run_keeps_what_it_read(tmp_path, monkeypatch):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"

    def mail_in_use(path):
        if path.name == "archive.pst":
            raise PermissionError(13, "Permission denied", str(path))
        return read_file(path)

    monkeypatch.setattr("document_manager.register.read_file", mail_in_use)
    with pytest.raises(Stopped):
        register_folder(docs, reg, stops_after(3))
    note, plan, mail = (docs / place for place in PLACES[:3])
    folder = str(docs.resolve())
    assert query(progress_of(reg), "SELECT kind, folder, path, size, changed, read_size,"
                                   " fingerprint, reason FROM run JOIN file ON run_id = run.id"
                                   " ORDER BY path") == [
        ("register", folder, "letters/note.txt", 22, note.stat().st_mtime_ns, 22,
         fingerprint(note), None),
        ("register", folder, "letters/plan.txt", 22, plan.stat().st_mtime_ns, 22,
         fingerprint(plan), None),
        ("register", folder, "mail/archive.pst", 22, mail.stat().st_mtime_ns, 22,
         None, IN_USE),
    ]


def test_lr2_progress_is_kept_once_a_minute_while_files_are_read(tmp_path, monkeypatch):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    clock = [0.0]
    monkeypatch.setattr("document_manager.progress.monotonic", lambda: clock[0])
    seen = []

    def counter(done, total):
        if total is None:
            return
        if done:
            seen.append(len(kept_places(reg)))  # what's kept once this many are read
        clock[0] += 25  # each file takes 25 seconds to read

    register_folder(docs, reg, counter)
    assert seen == [0, 0, 3, 3]  # kept once the first minute had passed, not after each file


def test_lr2_progress_is_kept_when_a_problem_stops_the_run(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"

    def counter(done, total):
        if (done, total) == (2, 4):
            docs.rename(tmp_path / "elsewhere")  # like a network drive that drops (BF-4)

    with pytest.raises(OSError, match="BF-4"):
        register_folder(docs, reg, counter)
    assert kept_places(reg) == PLACES[:2]
    assert not reg.exists()


def test_lr2_a_power_cut_loses_at_most_the_last_minute(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    cut = subprocess.Popen([sys.executable, "-c", POWER_CUT, "register", str(docs), str(reg)],
                           cwd=PROJECT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        give_up = time.monotonic() + 60
        while kept_places(reg) != PLACES[:2]:
            assert cut.poll() is None, "the run ended before the power cut"
            assert time.monotonic() < give_up, "the run never kept its first two files"
            time.sleep(0.05)
    finally:
        cut.kill()  # ended at once: nothing more runs in it
        cut.wait()
    assert not reg.exists()
    done = subprocess.run([sys.executable, "-m", "document_manager", "register", str(docs),
                           str(reg)], cwd=PROJECT, capture_output=True, text=True)
    assert done.returncode == 0, done.stderr
    assert re.fullmatch(r"Picked up the run that stopped at .+ UTC: 2 files kept, 0 read again\.\n"
                        rf"Registered {re.escape(str(docs))}\. Files listed: 4\.\n", done.stdout)
    fresh = tmp_path / "fresh.db"
    register_folder(docs, fresh)
    assert the_list(reg) == the_list(fresh)  # the same as a run that never stopped


def test_lr3_the_next_register_picks_up(tmp_path, monkeypatch):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    with pytest.raises(Stopped):
        register_folder(docs, reg, stops_after(3))
    (docs / "letters" / "note.txt").write_bytes(b"FAKE LETTERS/NOTE.TXT\n")  # the same size
    (docs / "letters" / "plan.txt").write_bytes(b"a longer plan, changed since\n")
    opened = reads(monkeypatch)
    files = register_folder(docs, reg)
    assert opened == ["note.txt", "plan.txt", "report.txt"]  # archive.pst is kept, not opened
    assert (files.picked_up.kept, files.picked_up.read_again) == (1, 2)
    fresh = tmp_path / "fresh.db"
    register_folder(docs, fresh)
    assert the_list(reg) == the_list(fresh)


def test_lr3_a_file_that_couldnt_be_read_before_is_read_again(tmp_path, monkeypatch):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    with held_open(docs / "mail" / "archive.pst"):
        with pytest.raises(Stopped):
            register_folder(docs, reg, stops_after(3))
    files = register_folder(docs, reg)
    assert (files.picked_up.kept, files.picked_up.read_again) == (2, 1)
    assert files.unread == []


def test_lr3_a_kept_file_isnt_opened(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    with pytest.raises(Stopped):
        register_folder(docs, reg, stops_after(3))
    with held_open(docs / "mail" / "archive.pst"):  # it couldn't be read now
        files = register_folder(docs, reg)
    assert files.unread == []  # kept from before, so it wasn't opened
    assert files.picked_up.kept == 3


def test_lr3_the_next_check_picks_up(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    (docs / "report.txt").write_bytes(b"a changed report\n")
    with pytest.raises(Stopped):
        check_folder(docs, reg, stops_after(2))
    found = check_folder(docs, reg)
    assert (found.changed, found.unchanged) == (["report.txt"], 3)
    assert (found.picked_up.kept, found.picked_up.read_again) == (2, 0)


def test_lr3_a_run_picks_up_only_from_its_own_kind(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    note = docs / "letters" / "note.txt"
    stale = Progress(reg, "register", docs.resolve())  # a Register's progress on the same folder
    stale.add("letters/note.txt", note.stat().st_size, note.stat().st_mtime_ns, 22,
              "0" * 64, None)
    stale.keep()
    found = check_folder(docs, reg)
    assert found == CheckResult(changed=[], new=[], missing=[], unchanged=4)  # not picked up
    assert kept_places(reg) == ["letters/note.txt"]  # and the Register's progress is still there


def test_lr3_a_run_picks_up_only_with_the_same_register_file(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    with pytest.raises(Stopped):
        register_folder(docs, reg, stops_after(2))
    other = tmp_path / "other.db"
    assert register_folder(docs, other).picked_up is None
    assert kept_places(reg) == PLACES[:2]


def test_lr4_register_says_it_picked_up(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    with pytest.raises(Stopped):
        run("register", str(docs), str(reg), stops_after(2))
    [(kept_at,)] = query(progress_of(reg), "SELECT kept_at FROM run")
    assert run("register", str(docs), str(reg)) == (0, (
        f"Picked up the run that stopped at {kept_at[:10]} {kept_at[11:16]} UTC:"
        " 2 files kept, 0 read again.\n"
        f"Registered {docs}. Files listed: 4."
    ))


def test_lr4_check_says_it_picked_up(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    with pytest.raises(Stopped):
        run("check", str(docs), str(reg), stops_after(1))
    code, words = run("check", str(docs), str(reg))
    assert code == 0
    assert re.fullmatch(r"Picked up the run that stopped at \d{4}-\d\d-\d\d \d\d:\d\d UTC:"
                        r" 1 file kept, 0 read again\.\nNothing has changed\. Files checked: 4\.",
                        words)


def test_lr5_progress_goes_when_its_run_finishes(tmp_path):
    a, b = make_docs(tmp_path / "a"), make_docs(tmp_path / "b")
    reg = tmp_path / "register.db"
    for docs in (a, b):
        with pytest.raises(Stopped):
            register_folder(docs, reg, stops_after(2))
    register_folder(a, reg)
    assert query(progress_of(reg), "SELECT folder FROM run") == [(str(b.resolve()),)]
    register_folder(b, reg)
    assert not progress_of(reg).exists()  # no unfinished run is left


def test_lr6_the_register_file_is_written_only_when_a_run_finishes(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    before = reg.read_bytes()
    with pytest.raises(Stopped):
        check_folder(docs, reg, stops_after(2))
    with pytest.raises(Stopped):
        register_folder(make_docs(tmp_path / "more"), reg, stops_after(2))
    assert reg.read_bytes() == before


def test_lr7_register_refuses_a_registered_folder_before_reading_any_file(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    stale = Progress(reg, "register", docs.resolve())  # left by a Register that can't finish
    stale.add("report.txt", 1, 1, 1, "0" * 64, None)
    stale.keep()
    told = []
    with pytest.raises(ValueError, match="RG-5"):
        register_folder(docs, reg, lambda done, total: told.append((done, total)))
    assert told == []  # not even a file found
    assert not progress_of(reg).exists()


def test_lr8_a_stopped_accept_or_summarize_keeps_nothing(tmp_path):
    docs = make_docs(tmp_path)
    for name in ("a.docx", "b.docx"):
        made = docx.Document()
        made.add_paragraph("A fake document for the tests.")
        made.save(docs / name)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    (docs / "report.txt").write_bytes(b"a changed report\n")
    with pytest.raises(Stopped):
        changes_to_accept(str(docs), str(reg), stops_after(2))  # Accept's first reading
    _, _, shown = changes_to_accept(str(docs), str(reg))
    with pytest.raises(Stopped):
        accept(str(docs), str(reg), shown, stops_after(2))  # and its second
    with pytest.raises(Stopped):
        run("summarize", str(docs), str(reg), stops_after(1))
    assert not progress_of(reg).exists()


def test_lr8_accept_doesnt_pick_up_a_stopped_check(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    (docs / "report.txt").write_bytes(b"a changed report\n")
    with pytest.raises(Stopped):
        check_folder(docs, reg, stops_after(2))
    code, words, found = changes_to_accept(str(docs), str(reg))
    assert (code, found.picked_up) == (0, None)
    assert not words.startswith("Picked up")
    accept_folder(docs, reg, found)
    assert kept_places(reg) == PLACES[:2]  # the Check's progress is left for the next Check


def test_lr9_only_the_register_file_and_the_progress_file_are_written(tmp_path):
    docs = make_docs(tmp_path)
    before = {path: path.read_bytes() for path in docs.rglob("*") if path.is_file()}
    with pytest.raises(Stopped):
        register_folder(docs, tmp_path / "register.db", stops_after(2))
    assert sorted(path.name for path in tmp_path.iterdir()) == ["docs", "register.db.progress"]
    assert {path: path.read_bytes() for path in docs.rglob("*") if path.is_file()} == before


def test_ctrl_c_stops_a_typed_command_and_keeps_what_it_read(tmp_path, monkeypatch, capsys):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    read = []

    def interrupted(f):
        read.append(f)
        if len(read) == 3:
            raise KeyboardInterrupt  # Ctrl+C
        return fingerprint_of(f)

    monkeypatch.setattr("document_manager.register.fingerprint_of", interrupted)
    try:
        code = main(["register", str(docs), str(reg)])
    except KeyboardInterrupt:
        pytest.fail("Ctrl+C wasn't caught")
    assert code == 1
    assert capsys.readouterr().out == (
        "Register stopped. The register file wasn't changed; the next Register on this"
        " folder picks up from here.\n")
    assert kept_places(reg) == PLACES[:2]


def test_ctrl_c_still_stops_summarize_as_before(tmp_path, monkeypatch):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    made = docx.Document()
    made.save(docs / "a.docx")

    def interrupted(path):
        raise KeyboardInterrupt  # Ctrl+C

    monkeypatch.setattr("document_manager.summaries.fingerprint", interrupted)
    with pytest.raises(KeyboardInterrupt):
        main(["summarize", str(docs), str(reg)])


def test_a_kept_file_counts_as_read(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    with pytest.raises(Stopped):
        register_folder(docs, reg, stops_after(2))
    told = []
    register_folder(docs, reg, lambda done, total: told.append((done, total)))
    assert told == FOUND + [(0, 4), (1, 4), (2, 4), (3, 4), (4, 4)]


def test_a_progress_file_of_an_unknown_version_is_refused(tmp_path):
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    con = sqlite3.connect(progress_of(reg))
    con.execute("PRAGMA user_version = 7")
    con.close()
    code, words = run("register", str(docs), str(reg))
    assert (code, words) == (
        1, f"Problem: Unknown progress file version: 7 ({progress_of(reg.resolve())})")
    assert not reg.exists()
