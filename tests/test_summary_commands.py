"""Tests for the summaries commands, spec/summaries.md v1.3 (SM-8).

Every Word document here is a fake, made by the test in a temporary folder.
"""

import re
import subprocess
import sys
from pathlib import Path

import docx

from document_manager.__main__ import main
from document_manager.register import register_folder
from document_manager.summaries import saved_summaries, summarize_folder
PROJECT = Path(__file__).resolve().parent.parent
OLD_FORMAT = bytes.fromhex("D0CF11E0A1B11AE1") + bytes(504)
OLD_REASON = "old Word file (.doc); save it as .docx in Word first"


def make_docs(tmp_path):
    """A small folder of fake documents."""
    docs = tmp_path / "docs"
    (docs / "letters").mkdir(parents=True)
    report = docx.Document()
    report.core_properties.title = "Report"
    report.add_heading("Summary", level=1)
    report.add_heading("Details", level=2)
    report.add_paragraph("Markets were calm.")
    report.save(docs / "report.docx")
    note = docx.Document()
    note.core_properties.title = ""
    note.add_paragraph("Dear Ernie.")
    note.save(docs / "letters" / "note.docx")
    (docs / "old.doc").write_bytes(OLD_FORMAT)
    return docs


def registered(tmp_path, capsys):
    """The fake documents, registered."""
    docs = make_docs(tmp_path)
    reg = tmp_path / "register.db"
    main(["register", str(docs), str(reg)])
    capsys.readouterr()
    return docs, reg


def test_sm8_summarize_says_what_it_did(tmp_path, capsys):
    docs, reg = registered(tmp_path, capsys)
    assert main(["summarize", str(docs), str(reg)]) == 0
    assert capsys.readouterr().out == (
        "Summarized: 2, unchanged: 0, couldn't read: 1.\n"
        f"  couldn't read  old.doc  {OLD_REASON}\n"
    )


def test_sm8_summaries_shows_each_document_and_the_couldnt_read_list(tmp_path, capsys):
    docs, reg = registered(tmp_path, capsys)
    main(["summarize", str(docs), str(reg)])
    capsys.readouterr()
    assert main(["summaries", str(docs), str(reg)]) == 0
    out = re.sub(r"\(summarized \d{4}-\d\d-\d\d \d\d:\d\d UTC\)", "(summarized TIME)",
                 capsys.readouterr().out)
    assert out == (
        f"Summaries for {docs}. Documents: 2, couldn't read: 1.\n"
        "\n"
        "letters/note.docx  (summarized TIME)\n"
        "  Title:    no title\n"
        "  Headings: none\n"
        "  Starts:   Dear Ernie.\n"
        "  Words:    2\n"
        "\n"
        "report.docx  (summarized TIME)\n"
        "  Title:    Report\n"
        "  Headings:\n"
        "    Summary\n"
        "      Details\n"
        "  Starts:   Summary Details Markets were calm.\n"
        "  Words:    5\n"
        "\n"
        "Couldn't read:\n"
        f"  old.doc  {OLD_REASON}\n"
    )


def test_sm8_no_summaries_yet(tmp_path, capsys):
    docs, reg = registered(tmp_path, capsys)
    assert main(["summaries", str(docs), str(reg)]) == 0
    assert capsys.readouterr().out == f"No summaries yet for {docs}. Run summarize first.\n"


def test_sm8_showing_only_reads_the_register_file(tmp_path, capsys):
    docs, reg = registered(tmp_path, capsys)
    main(["summarize", str(docs), str(reg)])
    before = (reg.read_bytes(), reg.stat().st_mtime_ns)
    main(["summaries", str(docs), str(reg)])
    assert (reg.read_bytes(), reg.stat().st_mtime_ns) == before


def test_sm8_unregistered_folder_is_a_problem_in_one_line(tmp_path, capsys):
    _, reg = registered(tmp_path, capsys)
    other = tmp_path / "other"
    other.mkdir()
    assert main(["summaries", str(other), str(reg)]) == 1
    out = capsys.readouterr().out
    assert out.startswith("Problem: ") and "SM-7" in out
    assert out.count("\n") == 1


def test_sm8_the_commands_work_as_written(tmp_path):
    docs = make_docs(tmp_path)
    odd = docx.Document()
    odd.core_properties.title = "Risk ✓ → watch"
    odd.add_paragraph("Rates ≥ 5%.")
    odd.save(docs / "odd.docx")
    reg = tmp_path / "register.db"
    command = [sys.executable, "-m", "document_manager"]
    runs = [
        subprocess.run([*command, name, str(docs), str(reg)], cwd=PROJECT,
                       capture_output=True, encoding="utf-8", errors="replace")
        for name in ("register", "summarize", "summaries")
    ]
    for run in runs:
        assert run.returncode == 0, run.stdout + run.stderr
    assert runs[1].stdout.startswith("Summarized: 3, unchanged: 0, couldn't read: 1.")
    assert "odd.docx" in runs[2].stdout and "Risk" in runs[2].stdout


def test_sm8_same_order_as_summarize_ignoring_capitals(tmp_path):
    docs = tmp_path / "docs"
    (docs / "letters").mkdir(parents=True)
    for place in ("Weekly.docx", "letters/note.docx", "apple.docx"):
        made = docx.Document()
        made.add_paragraph("Text.")
        made.save(docs / place)
    reg = tmp_path / "register.db"
    register_folder(docs, reg)
    order = ["apple.docx", "letters/note.docx", "Weekly.docx"]
    assert summarize_folder(docs, reg).summarized == order
    assert [s.place for s in saved_summaries(docs, reg)] == order
